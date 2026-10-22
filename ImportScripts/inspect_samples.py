"""Inspect sample rows across sheets to understand data types, hyperlinks, status codes, etc."""
import openpyxl
from collections import Counter

WB_PATH = r"C:\BenBotInventory\Builds.xlsx"

SHEETS = ["Builds", "Bulk 2026", "2025 Builds", "Bulk 2025", "2024 Builds", "Bulk 2024"]

wb = openpyxl.load_workbook(WB_PATH, data_only=True)

print("=========================================================")
print("STATUS LETTER FREQUENCIES (col 1) per sheet")
print("=========================================================")
for name in SHEETS:
    ws = wb[name]
    ctr = Counter()
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        ctr[repr(v)] += 1
    total_non_empty = sum(v for k, v in ctr.items() if k not in ("None", "' '"))
    print(f"{name}: total non-empty col1={total_non_empty}")
    for k, v in ctr.most_common():
        print(f"   {k}: {v}")
    print()

print("=========================================================")
print("SALE DATE VALUE SAMPLES per sheet (first 6 non-empty)")
print("=========================================================")
for name in SHEETS:
    ws = wb[name]
    hdr_row = 1
    # find Sale Date column
    sale_col = None
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=1, column=c).value
        if isinstance(h, str) and h.strip().lower() == "sale date":
            sale_col = c
            break
    print(f"{name}: Sale Date col = {sale_col}")
    if sale_col:
        shown = 0
        for r in range(2, ws.max_row + 1):
            v = ws.cell(row=r, column=sale_col).value
            if v is not None and v != "":
                print(f"   row {r}: type={type(v).__name__}  val={v!r}")
                shown += 1
                if shown >= 6:
                    break
    print()

print("=========================================================")
print("HYPERLINK EXAMPLES on price cells (Builds sheet)")
print("=========================================================")
ws = wb["Builds"]
found = 0
for r in range(2, min(ws.max_row + 1, 50)):
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row=r, column=c)
        if cell.hyperlink is not None:
            print(f"   row {r} col {c}: value={cell.value!r}  hyperlink={cell.hyperlink.target!r}")
            found += 1
            if found >= 10:
                break
    if found >= 10:
        break

print()
print("=========================================================")
print("EXAMPLE SOLD-TO CELLS from Builds sheet (col 34)")
print("=========================================================")
for r in range(2, 20):
    v = ws.cell(row=r, column=34).value
    if v:
        print(f"   row {r}: {v!r}")

print()
print("=========================================================")
print("EXAMPLE Note cells from Bulk 2026 (col 2)")
print("=========================================================")
ws = wb["Bulk 2026"]
shown = 0
for r in range(2, ws.max_row + 1):
    v = ws.cell(row=r, column=2).value
    if v:
        print(f"   row {r}: {v!r}")
        shown += 1
        if shown >= 6:
            break

print()
print("=========================================================")
print("How many non-empty rows (any part column has data)?")
print("=========================================================")
def part_cols(ws):
    """Part text cols: 3, 5, 7, 9, 11, 13, 15, 17, 19, 21, 23 (with matching price cols +1)."""
    return [3, 5, 7, 9, 11, 13, 15, 17, 19, 21, 23]

for name in SHEETS:
    ws = wb[name]
    count = 0
    for r in range(2, ws.max_row + 1):
        for c in part_cols(ws):
            if ws.cell(row=r, column=c).value not in (None, ""):
                count += 1
                break
    print(f"   {name}: {count} rows with at least one part")
