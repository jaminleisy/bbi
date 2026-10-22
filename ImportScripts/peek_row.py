"""Print full contents of specific rows for debugging."""
import openpyxl
import sys

WB_PATH = r"C:\BenBotInventory\Builds.xlsx"

wb = openpyxl.load_workbook(WB_PATH, data_only=True)

targets = [
    ("Builds",  [173, 174, 175, 176]),
    ("Builds",  [232, 235]),
    ("2025 Builds", [173, 174, 175]),
]

for sheet, rows in targets:
    ws = wb[sheet]
    print(f"=========== {sheet} ===========")
    for r in rows:
        print(f"--- row {r} ---")
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None and v != "":
                head = ws.cell(row=1, column=c).value
                print(f"  col {c:2d} ({head!r}): {v!r}")
        print()
