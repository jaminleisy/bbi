"""Discover unique seller codes and other patterns across all target sheets."""
import openpyxl
import re
from collections import Counter

WB_PATH = r"C:\BenBotInventory\Builds.xlsx"

SHEETS = ["Builds", "Bulk 2026", "2025 Builds", "Bulk 2025", "2024 Builds", "Bulk 2024"]

wb = openpyxl.load_workbook(WB_PATH, data_only=True)

part_text_cols = [3, 5, 7, 9, 11, 13, 15, 17, 19, 21, 23]

seller_ctr = Counter()
paren_examples = {}   # seller -> example
no_seller = []

for name in SHEETS:
    ws = wb[name]
    for r in range(2, ws.max_row + 1):
        for c in part_text_cols:
            v = ws.cell(row=r, column=c).value
            if not v or not isinstance(v, str):
                continue
            m = re.search(r"\(([^)]{1,60})\)", v)
            if m:
                seller = m.group(1).strip()
                seller_ctr[seller] += 1
                if seller not in paren_examples:
                    paren_examples[seller] = v.split("\n")[0][:100]
            else:
                if len(no_seller) < 40:
                    no_seller.append(f"{name} r{r} c{c}: {v.split(chr(10))[0][:100]}")

print("Distinct sellers/parenthetical codes:", len(seller_ctr))
print()
print("Top 50 seller codes with counts:")
for s, cnt in seller_ctr.most_common(50):
    print(f"  {cnt:5d}  '{s}'   e.g. {paren_examples[s]}")
print()
print("All rare seller codes (<= 3):")
for s, cnt in seller_ctr.most_common():
    if cnt <= 3:
        print(f"  {cnt:5d}  '{s}'   e.g. {paren_examples[s]}")

print()
print("=== Rows with parts but no parenthetical seller (first 40) ===")
for x in no_seller:
    print(f"  {x}")
